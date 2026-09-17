"""
CV Text Extraction Service
Supports: PDF (PyMuPDF), DOCX (python-docx)
"""
import os


def extract_text_from_file(file_path: str) -> str:
    """Extract plain text from PDF or DOCX file."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_from_docx(file_path)
    else:
        raise ValueError(f"Định dạng file không được hỗ trợ: {ext}")


def _extract_from_pdf(file_path: str) -> str:
    """Extract text from PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(file_path)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
        return "\n".join(text_parts).strip()
    except ImportError:
        raise RuntimeError("PyMuPDF chưa được cài đặt. Chạy: pip install pymupdf")
    except Exception as e:
        raise RuntimeError(f"Lỗi đọc file PDF: {e}")


def _extract_from_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        return "\n".join(paragraphs).strip()
    except ImportError:
        raise RuntimeError("python-docx chưa được cài đặt. Chạy: pip install python-docx")
    except Exception as e:
        raise RuntimeError(f"Lỗi đọc file DOCX: {e}")
